"""Build the WP-B generative chemistry competitor database workbook."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from competitors import COMPETITORS, DEALS, EXCLUDED, PRICING

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E5A")
THREAT_FILL = {"High": "F4C7C3", "Medium": "FCE8B2", "Low": "D9EAD3", "Complementor": "CFE2F3"}
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
ws = wb.active
ws.title = "Competitors"
cols = ["Company", "Tier", "HQ", "Product(s)", "Approach", "Business model", "Funding / financials",
        "Approx. capital raised (US$M)", "Key deals", "Validation evidence", "Threat to SyntheMol",
        "Why it matters", "Synthesis grounding (0-10)", "Commercial model (0 tools - 10 pipeline)", "Source"]
rows = [[c["name"], c["tier"], c["hq"], c["product"], c["approach"], c["model"], c["funding"],
         c["funding_usd_m"], c["deals"], c["evidence"], c["threat"], c["why"], c["x"], c["y"], c["src"]]
        for c in COMPETITORS]
sheet(ws, cols, rows, [30, 18, 22, 32, 45, 28, 38, 14, 38, 34, 12, 45, 12, 14, 45])
for row in ws.iter_rows(min_row=2):
    fill = THREAT_FILL.get(row[10].value)
    if fill:
        row[10].fill = PatternFill("solid", fgColor=fill)

sheet(wb.create_sheet("Pricing benchmarks"), ["Item", "Figure", "Source", "Confidence"], PRICING, [34, 60, 40, 16])
sheet(wb.create_sheet("Deal comps"), ["Deal", "Year", "Upfront (US$M)", "Headline value (US$M)", "Notes"],
      DEALS, [30, 8, 16, 20, 40])
sheet(wb.create_sheet("Excluded (biologics)"), ["Company", "Reason"], EXCLUDED, [24, 80])
notes = wb.create_sheet("Read me")
for line in [
    "WP-B Competitive Intelligence: generative chemistry (Dhruv). Research as of 30 Sep 2026.",
    "Scope: small-molecule generative chemistry. Generative biologics companies listed separately and excluded.",
    "Synthesis grounding and commercial-model scores are team judgements used for the positioning map, not published data.",
    "Capital raised is approximate and in US$; for public companies it is left blank or shows the IPO raise.",
    "Pricing rows marked 'verify' come from third-party estimates, not vendor price lists.",
]:
    notes.append([line])
notes.column_dimensions["A"].width = 110
wb.move_sheet("Read me", offset=-4)
wb.save("WP-B_GenChem_Competitor_Database.xlsx")
print("saved")
